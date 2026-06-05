"""One-shot generator for tests/fixtures/docx/sample_essay.docx.

Run once via: uv run python -m tests.extractors._make_fixture_docx
The output is committed; this script is here for reproducibility.
"""

from pathlib import Path

from docx import Document

OUT = Path(__file__).parent.parent / "fixtures" / "docx" / "sample_essay.docx"


def main() -> None:
    doc = Document()
    doc.add_heading("Body", level=1)
    doc.add_paragraph(
        "Climate change is well documented (Smith, 2020). Recent work "
        "extends this (Jones & Lee, 2021)."
    )
    doc.add_heading("References", level=1)
    doc.add_paragraph(
        "Jones, A., & Lee, B. (2021). New analyses of climate data. "
        "Climate Journal, 5(2), 100-120. https://doi.org/10.1234/cj.2021.05"
    )
    doc.add_paragraph("Smith, J. (2020). Foundations of climate research. Earth Press.")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
